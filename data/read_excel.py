import sys
import json
from openpyxl import load_workbook

if len(sys.argv) < 2:
    print(json.dumps({"error":"Necesita pasar el archivo de Excel"}, ensure_ascii=False))
    sys.exit(1)

xlsx = sys.argv[1]
sheet_names = sys.argv[2:]  # if empty, use all sheets

wb = load_workbook(xlsx, data_only=True)
if not sheet_names:
    sheet_names = wb.sheetnames

out = {}
for name in sheet_names:
    if name not in wb.sheetnames:
        out[name] = {"error": "sheet not found"}
        continue
    ws = wb[name]
    rows = list(ws.values)
    if not rows:
        out[name] = []
        continue
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    data = []
    for r in rows[1:]:
        row = {}
        for i, h in enumerate(headers):
            val = r[i] if i < len(r) else None
            if isinstance(val, str):
                val = val.strip()
            row[h] = val
        data.append(row)
    out[name] = data

print(json.dumps(out, ensure_ascii=False))
